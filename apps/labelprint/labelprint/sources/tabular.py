# -*- coding: utf-8 -*-
"""表格解析：.xls / .xlsx / .csv -> LabelRow 列表。

.xls 为 OLE2 老格式须用 xlrd；.xlsx 须用 openpyxl，两者不可混用
（xlrd 2.0 起已移除 xlsx 支持）。
"""
import csv
import io
import os

from labelprint.model import LabelRow

REQUIRED_COLUMN = "样品名称"

# 规范化列名 -> 字段名。键均为去空格后的小写形式。
_ALIASES = {
    "样品名称": "sample", "样品名": "sample", "sample": "sample",
    "浓度(ng/ul)": "conc", "浓度": "conc", "conc": "conc",
    "260/280": "ratio", "260-280": "ratio", "ratio": "ratio",
    "备注": "note", "note": "note",
}


class MissingColumnError(Exception):
    def __init__(self, message):
        super().__init__(message)
        self.message = message


def _norm(name) -> str:
    return str(name or "").strip().replace(" ", "").lower()


def _cell(value) -> str:
    """统一转成字符串；浮点数格式化到最多 6 位小数并去掉尾随零。

    Excel/表格里用公式算出来的数（比如 260/280 比值常见写法 =A2/B2）经常
    带二进制浮点表示误差，读出来是 1.8299999999999998 这种数，原样转字符串
    会整段贴到标签上，把这一行占满、后面的备注被挤没。这里只是消除浮点
    表示噪音，不是改数据本身——6 位小数早已超出这类实验室数据的有效精度。
    """
    if value is None:
        return ""
    if isinstance(value, float):
        s = "%.6f" % value
        s = s.rstrip("0").rstrip(".")
        return s if s else "0"
    return str(value).strip()


def _to_rows(table) -> list:
    """table 为二维列表，首行是表头。"""
    if not table:
        raise MissingColumnError("文件是空的，没有读到任何内容")
    header = [_norm(c) for c in table[0]]
    mapping = {}
    for idx, name in enumerate(header):
        field = _ALIASES.get(name)
        if field and field not in mapping:
            mapping[field] = idx
    if "sample" not in mapping:
        raise MissingColumnError(
            "表格里找不到「%s」列。读到的列名是：%s"
            % (REQUIRED_COLUMN, "、".join(str(c) for c in table[0]))
        )

    rows = []
    for raw in table[1:]:
        def pick(field):
            i = mapping.get(field)
            return _cell(raw[i]) if i is not None and i < len(raw) else ""
        row = LabelRow(sample=pick("sample"), conc=pick("conc"),
                       ratio=pick("ratio"), note=pick("note"))
        if not row.is_empty():
            rows.append(row)
    return rows


def _read_csv(data: bytes) -> list:
    for encoding in ("utf-8-sig", "gbk"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise MissingColumnError("CSV 编码无法识别，请另存为 UTF-8 或 GBK")
    return [r for r in csv.reader(io.StringIO(text))]


def _read_xlsx(data: bytes) -> list:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _read_xls(data: bytes) -> list:
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(i) for i in range(sheet.nrows)]


def parse(file_bytes: bytes, filename: str) -> list:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".csv":
        table = _read_csv(file_bytes)
    elif ext == ".xlsx":
        table = _read_xlsx(file_bytes)
    elif ext == ".xls":
        table = _read_xls(file_bytes)
    else:
        raise ValueError("只支持 .xls、.xlsx 和 .csv，收到的是 %s" % (ext or "无扩展名"))
    return _to_rows(table)
